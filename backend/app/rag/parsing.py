"""Парсинг документов в последовательность текстовых блоков с номером страницы и разделом."""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt", ".html", ".htm", ".xlsx", ".csv"}

# "1.", "1.2.", "2.3.1", "Глава 3", "Раздел II", "Статья 5", "Section 4"
_HEADING_RE = re.compile(
    r"^((\d+\.){1,4}\d*\.?\s+\S|(глава|раздел|статья|параграф|chapter|section|article)\s+[\dIVXLC]+)",
    re.IGNORECASE,
)


@dataclass
class Block:
    text: str
    page: int | None = None
    section: str | None = None
    is_heading: bool = False


def _clean(text: str) -> str:
    text = text.replace(" ", " ").replace("­", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _looks_like_heading(line: str) -> bool:
    line = line.strip()
    if not line or len(line) > 120 or line.endswith((".", ",", ";", ":")) and not _HEADING_RE.match(line):
        return False
    if _HEADING_RE.match(line) and len(line) < 120:
        return True
    letters = [c for c in line if c.isalpha()]
    return len(letters) >= 4 and all(c.isupper() for c in letters) and len(line) < 100


def _table_to_markdown(rows: list[list]) -> str:
    rows = [[_clean(str(c)) if c is not None else "" for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(c.replace("\n", " ") for c in r) + " |" for r in rows]
    lines.insert(1, "|" + " --- |" * width)
    return "\n".join(lines)


def _assign_sections(blocks: list[Block]) -> list[Block]:
    current = None
    for b in blocks:
        if b.is_heading:
            current = b.text.strip()[:150]
        b.section = b.section or current
    return blocks


# ---------------------------------------------------------------- PDF
def parse_pdf(path: Path) -> list[Block]:
    import pymupdf as fitz

    blocks: list[Block] = []
    with fitz.open(path) as doc:
        for page_no, page in enumerate(doc, start=1):
            items: list[tuple[float, Block]] = []
            table_rects = []
            try:
                for table in page.find_tables().tables:
                    md = _table_to_markdown(table.extract())
                    if md:
                        rect = fitz.Rect(table.bbox)
                        table_rects.append(rect)
                        items.append((rect.y0, Block(md, page_no)))
            except Exception:
                pass

            for x0, y0, x1, y1, text, _no, btype in page.get_text("blocks", sort=True):
                if btype != 0:
                    continue
                rect = fitz.Rect(x0, y0, x1, y1)
                if any(rect.intersects(t) for t in table_rects):
                    continue
                text = _clean(text)
                if not text:
                    continue
                # убираем колонтитулы с одним номером страницы
                if re.fullmatch(r"\d{1,4}", text):
                    continue
                first_line = text.split("\n", 1)[0]
                if _looks_like_heading(first_line) and "\n" in text:
                    items.append((y0, Block(first_line, page_no, is_heading=True)))
                    items.append((y0 + 0.01, Block(text.split("\n", 1)[1], page_no)))
                else:
                    items.append((y0, Block(text, page_no, is_heading=_looks_like_heading(text))))

            items.sort(key=lambda it: it[0])
            blocks.extend(b for _, b in items)
    return _assign_sections(blocks)


# ---------------------------------------------------------------- DOCX
def parse_docx(path: Path) -> list[Block]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    blocks: list[Block] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = _clean(para.text)
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            heading = style.startswith(("heading", "заголовок", "title")) or _looks_like_heading(text)
            blocks.append(Block(text, is_heading=heading))
        elif tag == "tbl":
            table = Table(child, document)
            rows = []
            for row in table.rows:
                cells, prev = [], None
                for cell in row.cells:  # объединённые ячейки повторяются — схлопываем
                    if cell._tc is not prev:
                        cells.append(cell.text)
                    prev = cell._tc
                rows.append(cells)
            md = _table_to_markdown(rows)
            if md:
                blocks.append(Block(md))
    return _assign_sections(blocks)


# ---------------------------------------------------------------- Markdown / TXT
def parse_markdown(path: Path) -> list[Block]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    blocks: list[Block] = []
    for para in re.split(r"\n\s*\n", text):
        para = _clean(para)
        if not para:
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", para)
        if m:
            heading, _, rest = para.partition("\n")
            blocks.append(Block(heading.lstrip("#").strip(), is_heading=True))
            if rest.strip():
                blocks.append(Block(rest.strip()))
        else:
            blocks.append(Block(para, is_heading=_looks_like_heading(para)))
    return _assign_sections(blocks)


# ---------------------------------------------------------------- HTML
def parse_html(path: Path) -> list[Block]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        tag.decompose()
    blocks: list[Block] = []
    for el in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "table", "pre"]):
        if el.name == "table":
            rows = [[c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])] for tr in el.find_all("tr")]
            md = _table_to_markdown(rows)
            if md:
                blocks.append(Block(md))
            continue
        if el.find_parent("table"):
            continue
        text = _clean(el.get_text(" ", strip=True))
        if text:
            blocks.append(Block(text, is_heading=el.name in {"h1", "h2", "h3", "h4"}))
    return _assign_sections(blocks)


# ---------------------------------------------------------------- XLSX / CSV
def parse_xlsx(path: Path) -> list[Block]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    blocks: list[Block] = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        blocks.append(Block(ws.title, is_heading=True))
        # большие таблицы режем по 40 строк, повторяя шапку
        header, body = rows[:1], rows[1:]
        for i in range(0, max(len(body), 1), 40):
            md = _table_to_markdown(header + body[i : i + 40])
            if md:
                blocks.append(Block(md))
    return _assign_sections(blocks)


def parse_csv(path: Path) -> list[Block]:
    import csv

    with path.open(encoding="utf-8", errors="ignore", newline="") as f:
        rows = list(csv.reader(f))
    header, body = rows[:1], rows[1:]
    return [Block(md) for i in range(0, max(len(body), 1), 40) if (md := _table_to_markdown(header + body[i : i + 40]))]


def parse_document(path: Path) -> list[Block]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext in {".md", ".markdown", ".txt"}:
        return parse_markdown(path)
    if ext in {".html", ".htm"}:
        return parse_html(path)
    if ext == ".xlsx":
        return parse_xlsx(path)
    if ext == ".csv":
        return parse_csv(path)
    raise ValueError(f"Неподдерживаемый формат: {ext}")
