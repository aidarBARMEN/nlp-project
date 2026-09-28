"""Text and spreadsheet formats accepted by the partner's upload interface."""

import csv
import re
from io import BytesIO, StringIO

from ingestion.models import Block, ParsedDocument


def parse_text(data: bytes, title: str) -> ParsedDocument:
    blocks: list[Block] = []
    section = None
    for paragraph in re.split(r"\n\s*\n", data.decode("utf-8-sig")):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if re.match(r"^#{1,6}\s+", paragraph):
            heading, _, rest = paragraph.partition("\n")
            section = heading.lstrip("#").strip()
            blocks.append(Block(text=section, kind="heading", section=section))
            if rest.strip():
                blocks.append(Block(text=rest.strip(), section=section))
        else:
            blocks.append(Block(text=paragraph, section=section))
    return ParsedDocument(title=title, blocks=blocks)


def table_blocks(rows, section=None) -> list[Block]:
    cleaned = [
        [
            str(cell).strip().replace("\n", " ").replace("|", "\\|") if cell is not None else ""
            for cell in row
        ]
        for row in rows
    ]
    cleaned = [row for row in cleaned if any(row)]
    if not cleaned:
        return []
    width = max(map(len, cleaned))
    cleaned = [row + [""] * (width - len(row)) for row in cleaned]
    header, body = cleaned[0], cleaned[1:]
    blocks = []
    for offset in range(0, max(len(body), 1), 40):
        batch = [header, ["---"] * width, *body[offset : offset + 40]]
        text = "\n".join("| " + " | ".join(row) + " |" for row in batch)
        blocks.append(Block(text=text, kind="table", section=section))
    return blocks


def parse_csv(data: bytes, title: str) -> ParsedDocument:
    stream = StringIO(data.decode("utf-8-sig"), newline="")
    return ParsedDocument(title=title, blocks=table_blocks(list(csv.reader(stream))))


def parse_xlsx(data: bytes, title: str) -> ParsedDocument:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        blocks = []
        for sheet in workbook.worksheets:
            blocks.extend(table_blocks(list(sheet.iter_rows(values_only=True)), sheet.title))
        return ParsedDocument(title=title, blocks=blocks)
    finally:
        workbook.close()
