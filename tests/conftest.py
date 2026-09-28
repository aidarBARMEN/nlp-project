from io import BytesIO

import pymupdf
import pytest
from docx import Document
from qdrant_client import QdrantClient

from ingestion.config import Settings
from ingestion.pipeline import Pipeline
from ingestion.services.indexer import Indexer


class TestEmbedder:
    """Deterministic test double, never selectable in production configuration."""

    dimension = 4

    def embed(self, texts):
        return [[1.0, float(len(t) % 7) / 7, 0.25, 0.5] for t in texts]

    def query(self, text):
        return self.embed([text])[0]


@pytest.fixture
def pipeline(tmp_path):
    settings = Settings(_env_file=None, data_dir=tmp_path / "store")
    indexer = Indexer(settings, QdrantClient(":memory:"))
    with Pipeline(settings, embedder=TestEmbedder(), indexer=indexer) as value:
        yield value


@pytest.fixture
def pdf_factory():
    def make(
        title="Academic Policy 2026",
        text="Students must register for exams. GPA and FX rules apply.",
        pages=2,
    ):
        with pymupdf.open() as document:
            document.set_metadata({"title": title})
            for i in range(pages):
                page = document.new_page()
                page.insert_text((72, 60), "Registration rules", fontsize=18)
                page.insert_text((72, 100), f"{text}\nPage {i + 1}", fontsize=11)
            return document.tobytes()

    return make


@pytest.fixture
def docx_bytes():
    document = Document()
    document.add_heading("Academic calendar 2026-2027", 0)
    document.add_heading("Examinations", 1)
    document.add_paragraph("Students register for exams in the student portal.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Grade", "Points"
    table.cell(1, 0).text, table.cell(1, 1).text = "A", "4.0"
    document.add_paragraph("After table: contact the registrar.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
